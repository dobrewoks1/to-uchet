"""Демо-данные «Учёт ТО»: вымышленный парк из 20 машин.

    python seed_demo.py [--data demo-data] [--xlsx ПУТЬ] [--force]

Скрипт строит синтетическую таблицу в формате импорта v8 (листы «Машины», «Журнал», «Шины», «Узлы»,
«Регламент ТО», «Настройки») и загружает её в базу <data>/to.db штатным импортом программы
(app.xlsx_io.import_v8) — тем же путём, что и кнопка «Загрузить из Excel». Книга сохраняется рядом
(<data>/demo-v8.xlsx) и, если задан --xlsx, ещё и по этому пути.

Все даты считаются от сегодняшнего дня, поэтому статусы всегда одинаковые: 4 машины с просроченным ТО,
4 — «скоро ТО», одна без внесённого ТО, одна без показаний, одна в ремонте, одна в простое, одна на ТО
сегодня, одна на консервации. Марки и модели обобщённые, госномера и VIN синтетические (регион «00»,
VIN начинается с DEMO), участки и фамилии вымышленные.

Потом: python run_dev.py --data demo-data  (или сразу: python run_dev.py --demo).
"""

import argparse
import io
import sys
from datetime import date, datetime, time, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app import db, xlsx_io  # noqa: E402
from app.logic import KIND_MARK  # noqa: E402

KM, MH = 'км', 'м/ч'

# Интервалы ТО — как в листе «Настройки» ниже: км-правила по маркам (N–P) и моточасы (B4/B5).
KM_RULES = [('КАМАЗ', 10000, 1000), ('SHACMAN', 10000, 1000), ('ШАХМАН', 10000, 1000), ('ШАКМАН', 10000, 1000),
            ('SITRAK', 15000, 1000), ('СИТРАК', 15000, 1000)]
INTERVAL_MH, WARN_MH, NORM_MH = 250, 50, 10

SITES = ['Северный', 'Южный', 'Мирный', 'Восточный', 'Западный', 'Промбаза']
TO_CYCLE = ['ТО-1', 'ТО-2', 'ТО-3']

# Парк. now — показание сейчас; rate — наработка в сутки; last — вид последнего ТО;
# left — сколько останется до следующего ТО (меньше нуля — просрочено); fresh — дней назад последняя отметка.
FLEET = [
    dict(no='101', model='КАМАЗ 6520', mtype='Самосвал', site='Северный', now=131200, rate=190,
         last='ТО-2', left=-1200, fresh=1),
    dict(no='102', model='КАМАЗ 6520', mtype='Самосвал', site='Северный', now=93900, rate=175,
         last='ТО-1', left=600, fresh=1),
    dict(no='103', model='КАМАЗ 65115', mtype='Самосвал', site='Южный', now=152300, rate=160,
         last='ТО-3', left=7700, fresh=3, autos=3),
    dict(no='104', model='КАМАЗ 65115', mtype='Самосвал', site='Южный', now=118400, rate=150,
         last='ТО-1', left=4100, fresh=1, open_repair=True),
    dict(no='105', model='КАМАЗ 65116 тягач', mtype='Тягач', site='Промбаза', now=86500, rate=140,
         last='ТО-2', left=10000, fresh=1, to_today=True),
    dict(no='106', model='SHACMAN X3000', mtype='Самосвал', site='Мирный', now=221000, rate=210,
         last='ТО-1', left=-11000, fresh=2),
    dict(no='107', model='SHACMAN X3000', mtype='Самосвал', site='Мирный', now=64800, rate=200,
         last='ТО-3', left=5600, fresh=2, autos=2),
    dict(no='108', model='SITRAK C7H', mtype='Самосвал', site='Восточный', now=74300, rate=220,
         last='ТО-2', left=700, fresh=1),
    dict(no='109', model='SITRAK C7H', mtype='Самосвал', site='Восточный', now=38200, rate=230,
         last=None, fresh=1),
    dict(no='110', model='КАМАЗ 6520', mtype='Самосвал', site='Промбаза', now=187500, rate=150,
         last='ТО-1', left=2500, fresh=120, status='На консервации'),
    dict(no='Э-01', model='Экскаватор гусеничный 22 т', mtype='Экскаватор', site='Северный', now=4487, rate=10,
         last='ТО-1', left=-27, fresh=1),
    dict(no='Э-02', model='Экскаватор колёсный 15 т', mtype='Экскаватор', site='Южный', now=2310.5, rate=9,
         last='ТО-3', left=120, fresh=1, repair_today=True),
    dict(no='Э-03', model='Экскаватор гусеничный 30 т', mtype='Экскаватор', site='Мирный', now=7795, rate=11,
         last='ТО-2', left=180, fresh=2, open_idle=True),
    dict(no='Э-04', model='Экскаватор-погрузчик 8 т', mtype='Экскаватор', site='Промбаза', now=960, rate=8.5,
         last='ТО-3', left=-12, fresh=1),
    dict(no='Б-01', model='Бульдозер 17 т', mtype='Бульдозер', site='Восточный', now=3168, rate=9.5,
         last='ТО-2', left=32, fresh=1),
    dict(no='Б-02', model='Бульдозер 25 т', mtype='Бульдозер', site='Западный', now=5420, rate=10,
         last='ТО-3', left=205, fresh=2),
    dict(no='П-01', model='Погрузчик фронтальный 3 т', mtype='Погрузчик', site='Промбаза', now=1890, rate=8,
         last='ТО-1', left=95, fresh=1),
    dict(no='П-02', model='Погрузчик фронтальный 5 т', mtype='Погрузчик', site='Западный', now=None),
    dict(no='К-01', model='Каток грунтовый 12 т', mtype='Каток', site='Западный', now=1245.5, rate=7,
         last='ТО-2', left=160, fresh=1),
    dict(no='Г-01', model='Автогрейдер 15 т', mtype='Грейдер', site='Южный', now=3911, rate=9,
         last='ТО-1', left=44, fresh=1),
]

TO_WORK = {
    KM: {'ТО-1': ('Замена масла ДВС и масляного фильтра, смазка шкворней и карданов',
                  'Масло 15W-40 — 28 л, фильтр масляный'),
         'ТО-2': ('Объём ТО-1 + топливные фильтры, проверка тормозов и рулевого',
                  'Масло 15W-40 — 28 л, фильтры масляный и топливный'),
         'ТО-3': ('Объём ТО-2 + масло КПП и мостов, воздушный фильтр',
                  'Масло 15W-40 — 28 л, 80W-90 — 30 л, комплект фильтров')},
    MH: {'ТО-1': ('Замена масла ДВС и фильтра, шприцовка шарниров',
                  'Масло 15W-40 — 24 л, фильтр масляный'),
         'ТО-2': ('Объём ТО-1 + топливные фильтры, проверка гидросистемы',
                  'Масло 15W-40 — 24 л, фильтры масляный и топливный'),
         'ТО-3': ('Объём ТО-2 + масло редукторов, фильтр гидравлики, воздушный фильтр',
                  'Масло 15W-40 — 24 л, гидравлическое — 20 л, комплект фильтров')},
}

# Закрытые ремонты и заметки в истории (по кругу на машины): дней назад, что делали, запчасти.
PAST_REPAIRS = [
    (9, 'Замена тормозных колодок задней оси', 'Колодки — 4 шт'),
    (17, 'Замена ремня генератора', 'Ремень генератора'),
    (24, 'Устранение течи гидроцилиндра ковша', 'Ремкомплект гидроцилиндра'),
    (13, 'Ремонт стартера', 'Бендикс стартера'),
    (31, 'Замена датчика давления масла', 'Датчик давления масла'),
]
NOTES = [
    (4, 'Шум в заднем мосту — проверить на ближайшем ТО'),
    (6, 'Замечаний нет, ходовая в норме'),
    (3, 'Подтекает сальник ступицы, наблюдаем'),
    (5, 'Проверен уровень масла в КПП — норма'),
]

TYRES = [
    # гар. №, № шины, тип, позиция, дата установки (дней назад), ресурс, пробег, примечание
    ('101', 'Ш-101-1', '12.00R20', 'Перед. левая', 420, 100000, 93000, 'с/н DEMO-0001'),
    ('101', 'Ш-101-2', '12.00R20', 'Перед. правая', 420, 100000, 91500, 'с/н DEMO-0002'),
    ('102', 'Ш-102-1', '12.00R20', 'Перед. левая', 300, 100000, 79000, 'с/н DEMO-0003'),
    ('106', 'Ш-106-1', '315/80R22.5', 'Зад. левая наружн.', 150, 120000, 40000, 'с/н DEMO-0004'),
    ('П-01', 'Ш-П01-1', '20.5-25', 'Перед. левая', 600, 6000, 5600, 'износ протектора'),
    ('К-01', 'Ш-К01-1', '23.5-25', 'Перед. левая', 200, 5000, 1200, ''),
]
UNITS = [
    # гар. №, узел, ресурс, показание машины при установке, рекомендация
    ('106', 'Двигатель', 270000, 0, 'Плановая диагностика при 250 000 км'),
    ('Э-01', 'Гидроцилиндр стрелы', 4800, 0, 'Ремкомплект заказать заранее'),
    ('Б-01', 'Ходовая часть', 8000, 500, ''),
    ('103', 'КПП', 250000, 20000, ''),
]

WHO_MARK, WHO_TO, WHO_REPAIR = 'Иванов', 'Петров', 'Сидоров'


def _interval(m):
    """Единица учёта, интервал и порог «скоро ТО» машины — по тем же правилам, что в программе."""
    words = ('экскаватор', 'бульдозер', 'каток', 'грейдер', 'погрузчик')
    text = f"{m['model']} {m['mtype']}".lower()
    if any(w in text for w in words):
        return MH, INTERVAL_MH, WARN_MH
    for keyword, interval, warn in KM_RULES:
        if keyword.lower() in m['model'].lower():
            return KM, interval, warn
    return MH, INTERVAL_MH, WARN_MH


def _round(value, unit):
    return int(round(value)) if unit == KM else round(value * 2) / 2      # км — целые, м/ч — до 0,5


def _prev_kind(kind):
    return TO_CYCLE[(TO_CYCLE.index(kind) - 1) % len(TO_CYCLE)]


def build_rows(today):
    """Строки листов книги v8 для парка FLEET на дату today: {лист: [{буква столбца: значение}]}."""
    def day(n):
        return datetime.combine(today - timedelta(days=n), time(0, 0))

    machines, journal, tyres, units = [], [], [], []
    for i, m in enumerate(FLEET):
        unit, interval, warn = _interval(m)
        n = i + 1
        letters = 'АВЕКМНОРСТХ'          # только буквы, которые бывают в госномерах; регион «00» не существует
        plate = (f'{letters[n % 11]}{100 + n:03d}{letters[(n * 3) % 11]}{letters[(n * 7) % 11]}00'
                 if unit == KM else f'{1000 + n:04d} ЭК 00')
        machine = {'A': n, 'B': m['model'], 'C': m['no'], 'D': plate, 'E': f'DEMO{n:013d}',
                   'G': m['site'], 'H': f'до 12.{2026 + n % 3}', 'K': m.get('status', 'В работе'), 'L': m['mtype']}
        machines.append(machine)
        if m['now'] is None:
            journal.append({'A': day(3), 'C': m['no'], 'D': 'Заметка',
                            'J': 'Пришла с другого участка, счётчик ещё не снят', 'N': WHO_MARK})
            continue

        now, rate, fresh = m['now'], m['rate'], m['fresh']

        def reading_at(days_ago):
            return _round(now - rate * (days_ago - fresh), unit)

        def days_at(value):
            """Сколько дней назад машина была на этом показании (по той же прямой, что и отметки)."""
            return int(round(fresh + (now - value) / rate))

        # ТО: последнее и предыдущее по кругу
        to_rows = []
        if m.get('last'):
            to_reading = now if m.get('to_today') else now - (interval - m['left'])
            to_days = 0 if m.get('to_today') else days_at(to_reading)
            prev_reading = to_reading - interval
            for kind, days_ago, value in ((_prev_kind(m['last']), days_at(prev_reading), prev_reading),
                                          (m['last'], to_days, to_reading)):
                if value <= 0:
                    continue
                work, parts = TO_WORK[unit][kind]
                row = {'A': day(days_ago), 'B': 'День', 'C': m['no'], 'D': kind, 'E': _round(value, unit),
                       'J': work, 'K': parts, 'N': WHO_TO}
                if days_ago == 0:
                    row.update({'F': time(8, 0), 'G': day(0), 'H': time(11, 30)})
                to_rows.append((days_ago, row))
            oldest = max(d for d, _ in to_rows) if to_rows else fresh + 70
        else:
            oldest = fresh + 70
        # показание при регистрации (без даты) — раньше всех отметок
        machine['F'] = max(_round(now - rate * (oldest + 14 - fresh), unit), 0)

        # Отметки показаний раз в неделю: от последней («fresh» дней назад) до самого старого ТО
        to_days_set = {d for d, _ in to_rows}
        days_ago = fresh
        while days_ago <= oldest + 7:
            if days_ago not in to_days_set:
                journal.append({'A': day(days_ago), 'C': m['no'], 'D': KIND_MARK,
                                'E': reading_at(days_ago), 'N': WHO_MARK})
            days_ago += 7
        journal.extend(row for _, row in to_rows)

        # «Накинуто за день» (как в выгрузке программы: показание пустое, «Кто» = «авто +N»)
        for d in range(m.get('autos', 0)):
            journal.append({'A': day(d), 'C': m['no'], 'D': KIND_MARK, 'N': f'авто +{_round(rate, unit)}'})

        # закрытый ремонт и заметка в истории
        r_days, r_work, r_parts = PAST_REPAIRS[i % len(PAST_REPAIRS)]
        if r_days < oldest and not m.get('status'):
            journal.append({'A': day(r_days), 'B': 'День', 'C': m['no'], 'D': 'Ремонт',
                            'E': reading_at(r_days), 'F': time(8, 30), 'G': day(r_days), 'H': time(14, 45),
                            'J': r_work, 'K': r_parts, 'N': WHO_REPAIR})
        n_days, n_text = NOTES[i % len(NOTES)]
        if m.get('status'):
            journal.append({'A': day(fresh - 1), 'C': m['no'], 'D': 'Заметка',
                            'J': 'Поставлена на консервацию до весны', 'N': WHO_MARK})
        elif not m.get('to_today') and not m.get('repair_today'):
            journal.append({'A': day(n_days), 'C': m['no'], 'D': 'Заметка', 'J': n_text, 'N': WHO_MARK})

        if m.get('open_repair'):
            journal.append({'A': day(1), 'B': 'День', 'C': m['no'], 'D': 'Ремонт', 'E': reading_at(1),
                            'F': time(8, 30), 'J': 'Течь масла из-под КПП, снята коробка',
                            'L': 'Сальник первичного вала КПП, прокладка поддона', 'N': WHO_REPAIR})
        if m.get('open_idle'):
            journal.append({'A': day(2), 'C': m['no'], 'D': 'Простой', 'F': time(7, 0),
                            'J': 'Ждёт трал для перебазировки на другой участок', 'N': WHO_MARK})
        if m.get('repair_today'):
            journal.append({'A': day(0), 'B': 'День', 'C': m['no'], 'D': 'Ремонт', 'E': _round(now, unit),
                            'F': time(8, 0), 'G': day(0), 'H': time(12, 30), 'J': 'Замена РВД рукояти',
                            'K': 'РВД 1 шт, масло гидравлическое — 10 л', 'N': WHO_REPAIR})

    # открытая заявка на запчасти по заметке («Заказано» = «нет»)
    journal.append({'A': day(5), 'C': '101', 'D': 'Заметка', 'J': 'Скол лобового стекла',
                    'L': 'Лобовое стекло', 'N': WHO_MARK, 'O': 'нет'})

    for no, tyre_no, ttype, pos, ago, resource, run, note in TYRES:
        tyres.append({'A': tyre_no, 'B': no, 'C': ttype, 'D': pos, 'E': day(ago), 'F': resource, 'G': run,
                      'K': note})
    for no, unit_name, resource, install, note in UNITS:
        units.append({'A': no, 'B': unit_name, 'C': resource, 'D': install, 'I': note})

    journal.sort(key=lambda r: (r['A'], r['C']))
    return {'Машины': machines, 'Журнал': journal, 'Шины': tyres, 'Узлы': units}


def build_workbook(today=None):
    """Книга xlsx в формате таблицы v8 (bytes)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    today = today or date.today()
    rows = build_rows(today)
    wb = Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)

    def sheet(title, caption, hint, header, data, widths):
        ws = wb.create_sheet(title)
        ws['A1'] = caption
        ws['A1'].font = Font(bold=True, size=14)
        ws['A2'] = hint
        for col, name in enumerate(header, 1):
            cell = ws.cell(row=4, column=col, value=name)
            cell.font = bold
        for r, row in enumerate(data, 5):
            for letter, value in row.items():
                if value is None or value == '':
                    continue
                cell = ws[f'{letter}{r}']
                cell.value = value
                if isinstance(value, datetime):
                    cell.number_format = 'DD.MM.YYYY'
                elif isinstance(value, time):
                    cell.number_format = 'HH:MM'
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = 'A5'
        return ws

    hint = 'Демо-данные: вымышленный парк, синтетические госномера и VIN. Формат — таблица учёта ТО v8.'
    sheet('Машины', 'Машины — заполняется один раз при запуске', hint,
          [name for name, _ in xlsx_io.MACHINE_COLUMNS[:12]], rows['Машины'],
          [5, 30, 11, 14, 20, 14, 14, 13, 14, 20, 16, 16])
    sheet('Журнал', 'Журнал — сюда пишем ВСЁ', hint,
          [name for name, _ in xlsx_io.JOURNAL_COLUMNS[:14]] + [xlsx_io.ORDERED_HEADER], rows['Журнал'],
          [12, 8, 10, 20, 13, 10, 12, 10, 12, 44, 34, 34, 20, 12, 12])
    sheet('Шины', 'Шины — учёт и износ', hint,
          [name for name, _ in xlsx_io.TYRE_COLUMNS[:11]], rows['Шины'],
          [12, 11, 14, 18, 13, 12, 14, 12, 8, 14, 20])
    sheet('Узлы', 'Узлы — ресурс, гарантия, износ', hint,
          [name for name, _ in xlsx_io.UNIT_COLUMNS[:9]], rows['Узлы'],
          [12, 22, 12, 14, 12, 12, 8, 14, 36])

    reg = [dict(zip('ABCDEFG', row)) for row in db.SEED_REGULATION]
    for row in reg:
        row['B'] = int(row['B'])
    sheet('Регламент ТО', 'Регламент ТО — что делается на каждом обслуживании', hint,
          ['Тип ТО', 'Периодичность, м/ч', 'Работы', 'Масло / жидкости', 'Фильтры', 'Прочие запчасти', 'Примечание'],
          reg, [8, 12, 60, 34, 34, 26, 40])

    ws = wb.create_sheet('Настройки')
    ws['A1'] = 'Настройки'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = 'Списки справа питают выпадающие меню. КМ-техника (N–P): слово в «Модели» → ТО в километрах.'
    heads = {'A': 'Параметр', 'B': 'Значение', 'D': 'События (для «Журнала»)', 'F': 'Узлы (для листа «Узлы»)',
             'H': 'Шины — тип/размер (для «Шины»)', 'J': 'Позиция колеса (для «Шины»)',
             'L': 'Участки / объекты (для «Машин»)', 'N': 'КМ-техника: слово в «Модели»', 'O': 'Интервал ТО, км',
             'P': 'Порог «Скоро ТО», км'}
    for letter, name in heads.items():
        ws[f'{letter}3'] = name
        ws[f'{letter}3'].font = bold
    ws['A4'], ws['B4'] = 'Интервал ТО, моточасов (техника без км-учёта)', INTERVAL_MH
    ws['A5'], ws['B5'] = 'Порог «Скоро ТО» (м/ч до ТО)', WARN_MH
    ws['A7'], ws['B7'] = 'Начисление моточасов в день (оценка, техника без км-учёта)', NORM_MH
    lists = (('D', db.SEED_LISTS['kinds']), ('F', db.SEED_LISTS['unit_names']), ('H', db.SEED_LISTS['tyre_types']),
             ('J', db.SEED_LISTS['positions']), ('L', SITES))
    for letter, values in lists:
        for r, value in enumerate(values, 4):
            ws[f'{letter}{r}'] = value
    for r, (keyword, interval, warn) in enumerate(KM_RULES, 4):
        ws[f'N{r}'], ws[f'O{r}'], ws[f'P{r}'] = keyword, interval, warn
    for letter, width in zip('ABCDEFGHIJKLMNOP', (52, 10, 2, 20, 2, 30, 2, 16, 2, 20, 2, 16, 2, 16, 12, 12)):
        ws.column_dimensions[letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def seed(data_dir, xlsx_path=None, force=False, today=None):
    """Создать <data_dir>/to.db с демо-парком. Возвращает отчёт импорта (dict)."""
    data_dir = Path(data_dir)
    data_dir.mkdir(parents=True, exist_ok=True)
    db_path = data_dir / 'to.db'
    if db_path.exists():
        if not force:
            raise SystemExit(f'База уже есть: {db_path}. Перезаписать демо-базу — с ключом --force.')
        for suffix in ('', '-wal', '-shm'):
            path = db_path.with_name(db_path.name + suffix)
            if path.exists():
                path.unlink()

    book = build_workbook(today)
    (data_dir / 'demo-v8.xlsx').write_bytes(book)
    if xlsx_path:
        Path(xlsx_path).parent.mkdir(parents=True, exist_ok=True)
        Path(xlsx_path).write_bytes(book)

    conn = db.open_db(db_path)
    try:
        with db.LOCK:
            return xlsx_io.import_v8(conn, book)
    finally:
        db.close_db(conn)


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding='utf-8')
        except (AttributeError, ValueError):
            pass
    parser = argparse.ArgumentParser(description='Учёт ТО — демо-данные (вымышленный парк из 20 машин)')
    parser.add_argument('--data', default='demo-data', help='папка данных (по умолчанию demo-data в проекте)')
    parser.add_argument('--xlsx', help='сохранить демо-таблицу v8 ещё и сюда (например, examples/demo-v8.xlsx)')
    parser.add_argument('--force', action='store_true', help='перезаписать существующую демо-базу')
    args = parser.parse_args(argv)
    data = Path(args.data)
    if not data.is_absolute():
        data = ROOT / data
    report = seed(data, args.xlsx, args.force)
    print(f'Демо-база: {data / "to.db"}')
    print(f'Машин: {report["machines_new"]}, записей журнала: {report["events_new"]}, шин: {report["tyres_new"]}, '
          f'узлов: {report["units_new"]}, открытых заявок на запчасти: {report["orders_open"]}')
    for warning in report['warnings']:
        print('  предупреждение:', warning)
    print(f'Запуск: python run_dev.py --data {args.data}  →  http://127.0.0.1:8765/')
    return 0


if __name__ == '__main__':
    sys.exit(main())
