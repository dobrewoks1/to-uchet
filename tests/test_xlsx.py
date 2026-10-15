"""Тесты импорта из таблицы v8 и экспорта в xlsx (SPEC, раздел 6 и 8).

Запуск из корня проекта:  .venv\\Scripts\\python.exe -m unittest tests.test_xlsx -v
Книга «как таблица v8» генерируется openpyxl прямо в тесте (минимальная, но со всеми оговорёнными случаями).
"""

import io
import unittest
from datetime import date, datetime, time, timedelta
from unittest import mock

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

from app import db, logic, xlsx_io
from app.logic import KIND_MARK

SERIAL_DATE = 46210      # серийная дата Excel
SERIAL_ISO = (date(1899, 12, 30) + timedelta(days=SERIAL_DATE)).isoformat()


def services_available():
    """Есть ли уже services.build_views (его пишет другой разработчик) — нужен только для листа «Парк»."""
    try:
        from app import services
    except ImportError:
        return False
    return hasattr(services, 'build_views')


def put(ws, row, values):
    """Записать значения в строку листа: {буква столбца: значение}."""
    for col, value in values.items():
        ws[f'{col}{row}'] = value


def make_v8():
    """Минимальная книга в формате таблицы v8 -> bytes."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ---- Журнал ----
    ws = wb.create_sheet('Журнал')
    ws['A1'] = 'Журнал — сюда пишем ВСЁ'
    ws['A3'] = '=IF(Настройки!$B$6="","▼ Первая пустая строка","")'
    put(ws, 4, dict(zip('ABCDEFGHIJKLMN', [
        'ДАТА *', 'Смена', 'МАШИНА *', 'СОБЫТИЕ *', 'МОТОЧАСЫ / КМ *', 'Заехала — ВРЕМЯ', 'Вышла — ДАТА',
        'Вышла — ВРЕМЯ', 'В ремонте ВСЕГО (само)', 'Что делалось / что сломалось',
        'Масло / запчасти (использовано)', 'Заказать запчасти (в закуп)', 'Замечания (стружка и т.п.)', 'Кто'])))
    for r in range(5, 3701):          # формула «В ремонте всего» протянута далеко вниз, как в настоящей таблице
        ws[f'I{r}'] = f'=IF(OR($A{r}="",$G{r}=""),"",$G{r}+$H{r}-$A{r}-$F{r})'
    # закрытый ремонт с «Заказать»; показание строкой, время заезда — timedelta
    put(ws, 5, {'A': datetime(2026, 6, 23), 'B': 'День', 'C': '5/104', 'D': 'Ремонт', 'E': '125 000',
                'F': timedelta(hours=9, minutes=30), 'G': datetime(2026, 6, 24), 'H': time(14, 0),
                'J': 'Замена сальника', 'K': 'сальник', 'L': 'сальник 2 шт', 'N': 'Иванов'})
    # открытый ремонт с «Заказать» (есть и в архиве — дубль)
    put(ws, 6, {'A': datetime(2026, 7, 1), 'B': 'ночь', 'C': '5/104', 'D': 'Ремонт', 'E': 126000,
                'F': '930', 'J': 'Течь масла', 'L': 'прокладка'})
    # формула в столбце даты -> даты нет -> строка пропускается с предупреждением
    put(ws, 7, {'A': '=TODAY()', 'C': '5/104', 'D': 'Заметка', 'J': 'формула вместо даты'})
    # ТО без показания, дата строкой
    put(ws, 8, {'A': '05.07.2026', 'C': '2С', 'D': 'ТО-1', 'J': 'ТО без показания'})
    # серийная дата, машина в другом регистре, дробное показание
    put(ws, 9, {'A': SERIAL_DATE, 'C': '3Э', 'D': 'Отметка моточасов', 'E': 1250.5})
    # машины нет в списке -> заглушка
    put(ws, 10, {'A': datetime(2026, 7, 3), 'C': '9/999', 'D': 'Заметка', 'M': 'неизвестная машина'})
    # нераспознанная дата
    put(ws, 11, {'A': 'вчера', 'C': '5/104', 'D': 'Заметка'})
    # нераспознанное время заезда; дата выхода раньше даты; время выхода — доля суток
    put(ws, 12, {'A': datetime(2026, 7, 4), 'C': '5/104', 'D': 'Простой', 'F': 'утром',
                 'G': datetime(2026, 7, 3), 'H': 0.625, 'J': 'нет водителя'})
    # вид с маленькой буквы; показание — формула массива; время выхода без даты выхода
    put(ws, 13, {'A': datetime(2026, 7, 5), 'C': '5/104', 'D': 'ремонт', 'F': time(8, 0), 'H': time(17, 0),
                 'J': 'стартер'})
    ws['E13'] = ArrayFormula('E13', '=SUM(A1:A2)')
    put(ws, 14, {'A': datetime(2026, 7, 6), 'D': 'Заметка', 'J': 'без машины'})
    put(ws, 15, {'A': datetime(2026, 7, 6), 'C': '5/104', 'J': 'без вида события'})
    # точная копия строки 5
    put(ws, 16, {'A': datetime(2026, 6, 23), 'B': 'День', 'C': '5/104', 'D': 'Ремонт', 'E': '125 000',
                 'F': timedelta(hours=9, minutes=30), 'G': datetime(2026, 6, 24), 'H': time(14, 0),
                 'J': 'Замена сальника', 'K': 'сальник', 'L': 'сальник 2 шт', 'N': 'Иванов'})
    # гаражный № русской «А», а в «Машинах» — латинской
    put(ws, 17, {'A': datetime(2026, 7, 7), 'C': '2А/317', 'D': 'Заметка', 'J': 'русская А'})
    # ТО с «Заказать»; показание с единицей; время числом 9.3 и строкой «17.05»
    put(ws, 18, {'A': datetime(2026, 7, 8), 'C': '5/104', 'D': 'ТО-2', 'E': '127 500 км', 'F': 9.3,
                 'G': datetime(2026, 7, 8), 'H': '17.05', 'J': 'ТО-2 по регламенту', 'L': 'фильтр'})
    # после 1481 пустой строки — ещё читается
    put(ws, 1500, {'A': datetime(2026, 7, 9), 'C': '2С', 'D': 'Заметка', 'M': 'далеко внизу'})
    # после 2099 пустых строк — уже нет
    put(ws, 3600, {'A': datetime(2026, 7, 10), 'C': '2С', 'D': 'Заметка', 'M': 'за пределом'})

    wb.create_sheet('Главная')['A1'] = 'Главная — весь парк одним взглядом (само)'

    # ---- Архив журнала ----
    ws = wb.create_sheet('Архив журнала')
    ws['A1'] = 'Архив журнала'
    put(ws, 5, {'A': datetime(2026, 5, 1), 'B': 'День', 'C': '5/104', 'D': 'ТО-1', 'E': 120000, 'K': 123321})
    put(ws, 6, {'A': datetime(2026, 7, 1), 'B': 'ночь', 'C': '5/104', 'D': 'Ремонт', 'E': 126000,
                'F': '930', 'J': 'Течь масла', 'L': 'прокладка'})

    # ---- Машины ----
    ws = wb.create_sheet('Машины')
    ws['A1'] = 'Машины — заполняется один раз при запуске'
    put(ws, 4, dict(zip('ABCDEFGHIJKL', [
        '№', 'Модель / тип техники', 'Гаражный №', 'Госномер', 'VIN', 'Моточасы / пробег при регистрации',
        'Участок / объект', 'Гарантийный срок', 'Фото таблички (ссылка)', 'Примечание', 'Статус машины',
        'Тип машины (для заказ-наряда)'])))
    for r in range(5, 61):
        ws[f'A{r}'] = f'=IF($B{r}="","",ROW()-4)'
    put(ws, 5, {'B': 'SHACMAN X3000', 'C': '5/104', 'D': 'C000TO00', 'E': 'TESTVIN0000000001', 'F': 47415,
                'G': 'СЕВЕРНЫЙ', 'H': datetime(2027, 1, 15), 'K': 'В работе', 'L': 'Самосвал'})
    put(ws, 6, {'B': 'SANY', 'C': '2С', 'F': '34 617', 'G': 'Южный', 'K': 'на консервации'})
    put(ws, 7, {'B': 'КАМАЗ без номера', 'D': 'А000АА00'})
    put(ws, 8, {'K': 'В работе'})                       # только статус — строка считается пустой
    put(ws, 9, {'B': 'Экскаватор SANY SY215', 'C': ' 3э ', 'E': 512, 'G': 'Карьер №5', 'L': 'экскаватор'})
    put(ws, 10, {'B': 'Дубль', 'C': '5/104 '})
    put(ws, 11, {'B': 'SITRAK', 'C': '2A/317', 'G': 'MИРНЫЙ', 'K': 'В работе', 'L': 'Самосвал'})   # латинские A и M

    # ---- Шины ----
    ws = wb.create_sheet('Шины')
    put(ws, 4, dict(zip('ABCDEFGHIJK', [
        '№ шины', 'Гаражный №', 'Тип / размер шины', 'Позиция', 'Дата установки', 'Ресурс, км / м·ч',
        'Пробег / наработка, км / м·ч', 'Остаток, км / м·ч', 'Износ', 'Статус', 'Примечание'])))
    for r in range(5, 51):
        ws[f'H{r}'] = f'=IF(OR($F{r}="",$G{r}=""),"",$F{r}-$G{r})'
        ws[f'J{r}'] = f'=IF($I{r}="","","🟢 Норма")'
    put(ws, 5, {'A': 22, 'B': '5/104', 'C': '12.00R20', 'D': 'перед. правая', 'E': datetime(2026, 6, 30),
                'F': 80000, 'G': '12 500', 'K': 'новая'})
    put(ws, 6, {'A': 'Ш-7', 'B': '2С', 'C': '11.00R20', 'D': 'Запасная', 'E': '30.06.26'})
    put(ws, 7, {'A': 22, 'B': '5/104', 'C': '12.00R20', 'D': 'Перед. правая', 'E': datetime(2026, 6, 30),
                'F': 80000, 'G': 12500, 'K': 'новая'})                               # дубль строки 5 целиком
    put(ws, 8, {'A': 5, 'C': '12.00R20'})                                             # нет машины

    # ---- Узлы ----
    ws = wb.create_sheet('Узлы')
    put(ws, 4, dict(zip('ABCDEFGHI', [
        'Машина (гар. №)', 'Узел', 'Ресурс / гарантия, м/ч', 'Моточасы машины при установке',
        'Наработка узла, м/ч', 'Остаток, м/ч', 'Износ', 'Статус', 'Рекомендация / запчасти / стоимость'])))
    for r in range(5, 51):
        ws[f'E{r}'] = f'=IF(OR($A{r}="",$D{r}=""),"",1)'
    put(ws, 5, {'A': '5/104', 'B': 'КПП', 'C': 5000, 'D': 0, 'I': 'масло 80W-90'})
    put(ws, 6, {'A': '3Э', 'B': 'Двигатель', 'C': '10 000', 'D': 1000.5})
    put(ws, 7, {'A': '5/104', 'B': 'КПП', 'C': 5000, 'D': 0, 'I': 'масло 80W-90'})    # дубль строки 5
    put(ws, 8, {'A': '5/104', 'C': 100})                                              # нет узла

    # ---- Регламент ТО ----
    ws = wb.create_sheet('Регламент ТО')
    put(ws, 4, dict(zip('ABCDEFG', ['Тип ТО', 'Периодичность, м/ч', 'Работы', 'Масло / жидкости', 'Фильтры',
                                    'Прочие запчасти', 'Примечание'])))
    put(ws, 5, {'A': 'ТО-1', 'B': 250, 'C': 'Работы первого ТО', 'D': 'Масло ДВС', 'E': 'Масляный фильтр',
                'F': '—', 'G': 'Базовое'})
    put(ws, 6, {'A': 'ТО-2', 'B': 500, 'C': 'Работы второго ТО', 'D': 'Масло ДВС', 'E': 'Два фильтра'})

    # ---- Настройки ----
    ws = wb.create_sheet('Настройки')
    put(ws, 3, {'A': 'Параметр', 'B': 'Значение', 'D': 'События', 'F': 'Узлы', 'H': 'Шины', 'J': 'Позиция',
                'L': 'Участки', 'N': 'КМ-техника: слово в «Модели»', 'O': 'Интервал ТО, км', 'P': 'Порог'})
    put(ws, 4, {'A': 'Интервал ТО, моточасов (техника без км-учёта)', 'B': 300, 'D': 'ТО-1', 'F': 'Двигатель',
                'H': '12.00R20', 'J': 'Перед. левая', 'L': 'Северный', 'N': 'SHACMAN', 'O': 12000, 'P': 1200})
    put(ws, 5, {'A': 'Порог «Скоро ТО» (м/ч до ТО)', 'B': 40, 'D': 'Регл. 2000', 'F': 'Новый узел',
                'L': 'Новый участок', 'N': 'SANY', 'O': 15000})
    put(ws, 6, {'A': 'GID листа «Журнал» (для кнопки перехода)', 'D': 'Ремонт', 'N': 'SITRAK', 'O': '15 000',
                'P': 1000})
    put(ws, 8, {'A': 'Начисление моточасов в день (оценка, техника без км-учёта)', 'B': 12})

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ParsersTest(unittest.TestCase):
    """Разбор значений ячеек (SPEC 6, «Разбор значений»)."""

    def test_cell(self):
        self.assertIsNone(xlsx_io._cell('=TODAY()'))
        self.assertIsNone(xlsx_io._cell(ArrayFormula('A1', '=SUM(B1:B2)')))
        self.assertIsNone(xlsx_io._cell('   '))
        self.assertIsNone(xlsx_io._cell('#N/A'))
        self.assertEqual(xlsx_io._cell('  текст '), 'текст')
        self.assertEqual(xlsx_io._cell(0), 0)
        self.assertEqual(xlsx_io._cell(time(9, 30)), time(9, 30))

    def test_date(self):
        p = xlsx_io._parse_date
        self.assertEqual(p(datetime(2026, 6, 23, 10, 5)), '2026-06-23')
        self.assertEqual(p(date(2026, 6, 23)), '2026-06-23')
        self.assertEqual(p('23.06.2026'), '2026-06-23')
        self.assertEqual(p('3.6.26'), '2026-06-03')
        self.assertEqual(p('2026-06-23'), '2026-06-23')
        self.assertEqual(p(SERIAL_DATE), SERIAL_ISO)
        self.assertEqual(p(46210.75), SERIAL_ISO)
        self.assertIsNone(p(15))
        self.assertIsNone(p('вчера'))
        self.assertIsNone(p('31.02.2026'))
        self.assertIsNone(p(True))

    def test_time(self):
        p = xlsx_io._parse_time
        self.assertEqual(p(time(9, 30)), '09:30')
        self.assertEqual(p(datetime(1899, 12, 30, 15, 0)), '15:00')
        self.assertEqual(p(timedelta(hours=9, minutes=30)), '09:30')
        self.assertEqual(p(0.625), '15:00')
        self.assertEqual(p(0), '00:00')
        self.assertEqual(p('9:30'), '09:30')
        self.assertEqual(p('09:30:00'), '09:30')
        self.assertEqual(p('930'), '09:30')
        self.assertEqual(p('1430'), '14:30')
        self.assertEqual(p('9.30'), '09:30')
        self.assertEqual(p(9.3), '09:30')
        self.assertEqual(p(9.05), '09:05')
        self.assertEqual(p(930), '09:30')
        self.assertEqual(p('9'), '09:00')
        self.assertEqual(p(14), '14:00')
        self.assertIsNone(p('утром'))
        self.assertIsNone(p('25:00'))
        self.assertIsNone(p('9:75'))

    def test_number(self):
        p = xlsx_io._parse_number
        self.assertEqual(p(125000), 125000.0)
        self.assertEqual(p('125 000'), 125000.0)
        self.assertEqual(p('125\xa0000'), 125000.0)
        self.assertEqual(p('1250,5'), 1250.5)
        self.assertEqual(p('127 500 км'), 127500.0)
        self.assertEqual(p('1250,5 м/ч'), 1250.5)
        self.assertIsNone(p('много'))
        self.assertIsNone(p(True))
        self.assertIsNone(p(datetime(2026, 1, 1)))

    def test_text(self):
        t = xlsx_io._text
        self.assertEqual(t(123321), '123321')
        self.assertEqual(t(512.0), '512')
        self.assertEqual(t(datetime(2027, 1, 15)), '15.01.2027')
        self.assertEqual(t(None), '')


class ImportTest(unittest.TestCase):
    """Импорт минимальной книги v8."""

    @classmethod
    def setUpClass(cls):
        cls.book = make_v8()

    def setUp(self):
        self.conn = db.open_db(':memory:')
        self.addCleanup(self.conn.close)

    def one(self, sql, *args):
        return self.conn.execute(sql, args).fetchone()

    def machine_id(self, garage_no):
        return self.one('SELECT id FROM machines WHERE garage_no=?', garage_no)['id']

    def event(self, garage_no, day, kind=None):
        sql = 'SELECT * FROM events WHERE machine_id=? AND date=?' + (' AND kind=?' if kind else '')
        args = (self.machine_id(garage_no), day) + ((kind,) if kind else ())
        rows = self.conn.execute(sql, args).fetchall()
        self.assertEqual(len(rows), 1, f'{garage_no} {day} {kind}: ожидалось одно событие, найдено {len(rows)}')
        return rows[0]

    def assertWarning(self, result, *parts):
        """В предупреждениях есть строка, содержащая все подстроки parts."""
        for w in result['warnings']:
            if all(p in w for p in parts):
                return
        self.fail(f'нет предупреждения с {parts}; есть:\n' + '\n'.join(result['warnings']))

    def test_counts(self):
        r = xlsx_io.import_v8(self.conn, self.book)
        self.assertIs(r['ok'], True)
        self.assertIs(r['dry_run'], False)
        self.assertEqual(r['machines_new'], 5)          # 4 из листа «Машины» + заглушка 9/999
        self.assertEqual(r['machines_updated'], 0)
        self.assertEqual(r['events_new'], 11)           # 2 из архива + 9 из журнала
        self.assertEqual(r['events_skipped'], 2)        # строка 6 (есть в архиве) и строка 16 (копия строки 5)
        self.assertEqual(r['tyres_new'], 2)
        self.assertEqual(r['units_new'], 2)
        self.assertEqual(r['orders_open'], 1)
        self.assertEqual(r['settings_updated'], 8)      # 3 настройки + вид + узел + участок + правила + регламент
        self.assertEqual(r['warnings_total'], len(r['warnings']))
        self.assertFalse(self.conn.in_transaction)
        self.assertEqual(self.one('SELECT COUNT(*) FROM machines')[0], 5)
        self.assertEqual(self.one('SELECT COUNT(*) FROM events')[0], 11)
        self.assertEqual(self.one("SELECT COUNT(*) FROM events WHERE source<>'import'")[0], 0)

    def test_machines(self):
        r = xlsx_io.import_v8(self.conn, self.book)
        m = self.one('SELECT * FROM machines WHERE garage_no=?', '5/104')
        self.assertEqual(m['model'], 'SHACMAN X3000')
        self.assertEqual(m['plate'], 'C000TO00')
        self.assertEqual(m['vin'], 'TESTVIN0000000001')
        self.assertEqual(m['reg_reading'], 47415)
        self.assertIsNone(m['reg_date'])
        self.assertEqual(m['site'], 'Северный')                   # «СЕВЕРНЫЙ» -> написание из справочника
        self.assertEqual(m['warranty'], '15.01.2027')
        self.assertEqual(m['status'], 'В работе')
        self.assertEqual(m['mtype'], 'Самосвал')
        self.assertIsNotNone(m['created_at'])

        m = self.one('SELECT * FROM machines WHERE garage_no=?', '2С')
        self.assertEqual(m['reg_reading'], 34617)                 # «34 617» строкой
        self.assertEqual(m['status'], 'На консервации')           # «на консервации» -> как в списке статусов
        self.assertEqual(m['mtype'], '')

        m = self.one('SELECT * FROM machines WHERE garage_no=?', '3э')       # пробелы по краям убраны
        self.assertEqual(m['vin'], '512')                         # VIN числом -> текст
        self.assertEqual(m['status'], 'В работе')                 # пустой статус
        self.assertEqual(m['mtype'], 'Экскаватор')
        self.assertEqual(m['site'], 'Карьер №5')

        self.assertEqual(self.one('SELECT site FROM machines WHERE garage_no=?', '2A/317')['site'], 'Мирный')

        stub = self.one('SELECT * FROM machines WHERE garage_no=?', '9/999')
        self.assertEqual((stub['model'], stub['status'], stub['reg_reading']), ('', 'В работе', None))

        # порядок машин — как в таблице
        order = [row['garage_no'] for row in self.conn.execute('SELECT garage_no FROM machines ORDER BY pos, id')]
        self.assertEqual(order, ['5/104', '2С', '3э', '2A/317', '9/999'])

        self.assertWarning(r, 'Машины, строка 7', 'нет гаражного №')
        self.assertWarning(r, 'Машины, строка 10', 'повторяется')
        self.assertWarning(r, '2С', 'проверьте единицу учёта (км или моточасы)')
        self.assertFalse(any('строка 8' in w and 'Машины' in w for w in r['warnings']))
        self.assertFalse(any('3э' in w and 'единицу учёта' in w for w in r['warnings']))
        # участок, которого не было в справочнике, добавлен в него
        self.assertIn('Карьер №5', db.get_lists(self.conn)['sites'])

    def test_events(self):
        r = xlsx_io.import_v8(self.conn, self.book)

        ev = self.event('5/104', '2026-06-23')
        self.assertEqual(ev['kind'], 'Ремонт')
        self.assertEqual(ev['shift'], 'День')
        self.assertEqual(ev['reading'], 125000)                   # «125 000» строкой
        self.assertEqual(ev['time_in'], '09:30')                  # timedelta
        self.assertEqual(ev['date_out'], '2026-06-24')
        self.assertEqual(ev['time_out'], '14:00')
        self.assertEqual((ev['work'], ev['parts_used'], ev['parts_order'], ev['who']),
                         ('Замена сальника', 'сальник', 'сальник 2 шт', 'Иванов'))
        self.assertEqual(ev['source'], 'import')
        self.assertEqual(ev['reading_estimated'], 0)
        self.assertEqual(ev['order_done'], 1)                     # закрытый ремонт — заявка считается выполненной

        ev = self.event('5/104', '2026-07-01')
        self.assertEqual(ev['shift'], 'Ночь')                     # «ночь» -> написание из справочника
        self.assertEqual(ev['time_in'], '09:30')                  # «930»
        self.assertIsNone(ev['date_out'])
        self.assertEqual(ev['order_done'], 0)                     # открытый ремонт — заявка открыта

        ev = self.event('2С', '2026-07-05')
        self.assertEqual(ev['kind'], 'ТО-1')
        self.assertIsNone(ev['reading'])
        self.assertWarning(r, 'Журнал, строка 8', 'ТО-1 без показания')

        ev = self.event('3э', SERIAL_ISO)                         # серийная дата, «3Э» = «3э»
        self.assertEqual(ev['kind'], KIND_MARK)
        self.assertEqual(ev['reading'], 1250.5)

        ev = self.event('9/999', '2026-07-03')
        self.assertEqual(ev['remarks'], 'неизвестная машина')
        self.assertWarning(r, 'Журнал, строка 10', '9/999', 'заглушка')

        ev = self.event('5/104', '2026-07-04')
        self.assertEqual(ev['kind'], 'Простой')
        self.assertIsNone(ev['time_in'])                          # «утром» не распознано
        self.assertEqual(ev['date_out'], '2026-07-03')            # импортировано как есть
        self.assertEqual(ev['time_out'], '15:00')                 # доля суток 0.625
        self.assertWarning(r, 'Журнал, строка 12', 'время заезда', 'утром')
        self.assertWarning(r, 'Журнал, строка 12', 'раньше даты заезда')

        ev = self.event('5/104', '2026-07-05')
        self.assertEqual(ev['kind'], 'Ремонт')                    # «ремонт» -> «Ремонт»
        self.assertIsNone(ev['reading'])                          # формула массива — пусто
        self.assertEqual((ev['time_in'], ev['date_out'], ev['time_out']), ('08:00', None, '17:00'))
        self.assertWarning(r, 'Журнал, строка 13', 'без даты выхода')

        ev = self.event('2A/317', '2026-07-07')                   # русская «А» отнесена к машине с латинской
        self.assertEqual(ev['work'], 'русская А')
        self.assertWarning(r, 'Журнал, строка 17', '2А/317', '2A/317')

        ev = self.event('5/104', '2026-07-08')
        self.assertEqual(ev['kind'], 'ТО-2')
        self.assertEqual(ev['reading'], 127500)                   # «127 500 км»
        self.assertEqual((ev['time_in'], ev['time_out']), ('09:30', '17:05'))
        self.assertEqual(ev['order_done'], 1)

        ev = self.event('5/104', '2026-05-01')                    # из архива
        self.assertEqual((ev['kind'], ev['reading'], ev['parts_used']), ('ТО-1', 120000, '123321'))

        self.assertEqual(self.event('2С', '2026-07-09')['remarks'], 'далеко внизу')
        # после 2000 пустых строк чтение листа прекращается
        self.assertEqual(self.one("SELECT COUNT(*) FROM events WHERE remarks='за пределом'")[0], 0)

        self.assertWarning(r, 'Журнал, строка 7', 'не заполнено', 'дата')          # формула вместо даты
        self.assertWarning(r, 'Журнал, строка 11', 'вчера', 'не распознана')
        self.assertWarning(r, 'Журнал, строка 14', 'не заполнено', 'машина')
        self.assertWarning(r, 'Журнал, строка 15', 'не заполнено', 'событие')
        # архивные события вставлены раньше журнальных
        first = self.one('SELECT date FROM events ORDER BY id LIMIT 1')['date']
        self.assertEqual(first, '2026-05-01')

    def test_order_done(self):
        r = xlsx_io.import_v8(self.conn, self.book)
        rows = {(row['date'], row['kind']): row['order_done'] for row in self.conn.execute(
            "SELECT date, kind, order_done FROM events WHERE parts_order<>''")}
        self.assertEqual(rows, {('2026-06-23', 'Ремонт'): 1, ('2026-07-01', 'Ремонт'): 0, ('2026-07-08', 'ТО-2'): 1})
        self.assertEqual(r['orders_open'], 1)
        self.assertEqual(self.one("SELECT COUNT(*) FROM events WHERE parts_order='' AND order_done=1")[0], 0)

    def test_tyres_units(self):
        r = xlsx_io.import_v8(self.conn, self.book)
        tyres = self.conn.execute('SELECT t.*, m.garage_no FROM tyres t JOIN machines m ON m.id=t.machine_id '
                                  'ORDER BY t.id').fetchall()
        self.assertEqual(len(tyres), 2)
        t = tyres[0]
        self.assertEqual((t['garage_no'], t['tyre_no'], t['ttype'], t['position'], t['install_date'],
                          t['resource'], t['run'], t['install_reading'], t['note']),
                         ('5/104', '22', '12.00R20', 'Перед. правая', '2026-06-30', 80000, 12500, None, 'новая'))
        t = tyres[1]
        self.assertEqual((t['garage_no'], t['tyre_no'], t['install_date'], t['resource']),
                         ('2С', 'Ш-7', '2026-06-30', None))

        units = self.conn.execute('SELECT u.*, m.garage_no FROM units u JOIN machines m ON m.id=u.machine_id '
                                  'ORDER BY u.id').fetchall()
        self.assertEqual(len(units), 2)
        u = units[0]
        self.assertEqual((u['garage_no'], u['unit'], u['resource'], u['install_reading'], u['note']),
                         ('5/104', 'КПП', 5000, 0, 'масло 80W-90'))
        u = units[1]
        self.assertEqual((u['garage_no'], u['unit'], u['resource'], u['install_reading']),
                         ('3э', 'Двигатель', 10000, 1000.5))
        self.assertWarning(r, 'Шины, строка 8', 'нет гаражного №')
        self.assertWarning(r, 'Узлы, строка 8', 'узел')
        self.assertWarning(r, 'Шины, строка 7', 'совпадает со строкой 5')
        self.assertWarning(r, 'Узлы, строка 7', 'совпадает со строкой 5')
        self.assertEqual((r['tyres_skipped'], r['units_skipped']), (1, 1))

    def test_settings_lists_rules_regulation(self):
        # суточная норма совпадающего слова должна сохраниться
        self.conn.execute("UPDATE rules SET daily_norm=120 WHERE keyword='SHACMAN'")
        self.conn.commit()
        xlsx_io.import_v8(self.conn, self.book)

        s = db.get_settings(self.conn)
        self.assertEqual((s['default_interval_mh'], s['default_warn_mh'], s['default_norm_mh']), (300, 40, 12))
        self.assertEqual(s['default_interval_km'], 10000)         # остальное не тронуто

        lists = db.get_lists(self.conn)
        self.assertIn('Регл. 2000', lists['kinds'])
        self.assertEqual(lists['kinds'].count('Ремонт'), 1)       # существующие значения не задвоились
        self.assertIn('Новый узел', lists['unit_names'])
        self.assertIn('Новый участок', lists['sites'])
        self.assertEqual(lists['sites'][0], 'Северный')           # прежние значения остались на месте

        rules = [(r['keyword'], r['unit'], r['interval'], r['warn'], r['daily_norm'])
                 for r in db.get_rules(self.conn)]
        self.assertEqual(rules, [
            ('SHACMAN', 'км', 12000, 1200, 120),
            ('SANY', 'км', 15000, None, None),
            ('SITRAK', 'км', 15000, 1000, None),
            ('САНИ', 'км', 15000, 1000, None),
            ('VOLVO', 'км', 15000, 1000, None),
            ('ВОЛЬВО', 'км', 15000, 1000, None),
        ])

        reg = db.get_regulation(self.conn)
        self.assertEqual([(x['to_kind'], x['periodicity'], x['works']) for x in reg],
                         [('ТО-1', '250', 'Работы первого ТО'), ('ТО-2', '500', 'Работы второго ТО')])
        self.assertEqual(reg[1]['parts'], '')

    def test_tz_km_rules_added_when_missing_in_table(self):
        """Таблица v10: SANY убран из столбцов N–P — правило SANY 15 000 км (ТЗ, R3) всё равно есть после импорта,
        самосвал SANY считается в км, экскаватор SANY остаётся на моточасах."""
        wb = openpyxl.load_workbook(io.BytesIO(self.book))
        ws = wb['Настройки']
        ws['N5'] = None
        ws['O5'] = None
        buf = io.BytesIO()
        wb.save(buf)
        xlsx_io.import_v8(self.conn, buf.getvalue())
        rules = {r['keyword']: (r['unit'], r['interval']) for r in db.get_rules(self.conn)}
        for word in ('SANY', 'САНИ', 'VOLVO', 'ВОЛЬВО'):
            self.assertEqual(rules.get(word), ('км', 15000), word)
        self.assertEqual(rules['SHACMAN'], ('км', 12000))
        settings = db.get_settings(self.conn)
        rule_list = db.get_rules(self.conn)
        truck = logic.machine_params({'model': 'SANY', 'mtype': 'Самосвал'}, rule_list, settings)
        self.assertEqual((truck['unit'], truck['interval']), ('км', 15000))
        excavator = logic.machine_params({'model': 'Экскаватор SANY 330', 'mtype': 'Экскаватор'}, rule_list, settings)
        self.assertEqual(excavator['unit'], 'м/ч')

    def test_non_km_rules_kept(self):
        """Правила не в км (добавленные пользователем) импорт не трогает."""
        rules = db.get_rules(self.conn) + [{'keyword': 'CASE', 'unit': 'м/ч', 'interval': 500, 'warn': 50,
                                           'daily_norm': 10}]
        db.set_rules(self.conn, rules)
        xlsx_io.import_v8(self.conn, self.book)
        last = db.get_rules(self.conn)[-1]
        self.assertEqual((last['keyword'], last['unit'], last['interval']), ('CASE', 'м/ч', 500))

    def test_dry_run(self):
        before = (db.get_settings(self.conn), db.get_lists(self.conn), db.get_rules(self.conn),
                  db.get_regulation(self.conn))
        dry = xlsx_io.import_v8(self.conn, self.book, dry_run=True)
        self.assertIs(dry['dry_run'], True)
        self.assertEqual(dry['machines_new'], 5)
        self.assertEqual(dry['events_new'], 11)
        self.assertFalse(self.conn.in_transaction)
        for table in ('machines', 'events', 'tyres', 'units'):
            self.assertEqual(self.one(f'SELECT COUNT(*) FROM {table}')[0], 0, table)
        after = (db.get_settings(self.conn), db.get_lists(self.conn), db.get_rules(self.conn),
                 db.get_regulation(self.conn))
        self.assertEqual(before, after)

        # настоящий импорт после dry_run даёт тот же отчёт
        real = xlsx_io.import_v8(self.conn, self.book)
        self.assertEqual({k: v for k, v in dry.items() if k != 'dry_run'},
                         {k: v for k, v in real.items() if k != 'dry_run'})

    def test_repeat_import_no_duplicates(self):
        xlsx_io.import_v8(self.conn, self.book)
        counts = [self.one(f'SELECT COUNT(*) FROM {t}')[0] for t in ('machines', 'events', 'tyres', 'units')]
        lists = db.get_lists(self.conn)
        r = xlsx_io.import_v8(self.conn, self.book)
        self.assertEqual((r['machines_new'], r['machines_updated'], r['events_new'], r['tyres_new'],
                          r['units_new'], r['settings_updated'], r['orders_open']), (0, 0, 0, 0, 0, 0, 0))
        self.assertEqual(r['events_skipped'], 13)
        self.assertEqual(counts,
                         [self.one(f'SELECT COUNT(*) FROM {t}')[0] for t in ('machines', 'events', 'tyres', 'units')])
        self.assertEqual(lists, db.get_lists(self.conn))

    def test_update_existing_machine(self):
        """Upsert: заполненные в таблице поля обновляются, пустые не затирают, reg_date не трогается."""
        self.conn.execute(
            "INSERT INTO machines(pos, garage_no, model, plate, note, status, reg_reading, reg_date, norm_override) "
            "VALUES(1, '5/104', 'Старая модель', 'X000XX00', 'моя заметка', 'В работе', 100, '2026-01-01', 90)")
        self.conn.commit()
        r = xlsx_io.import_v8(self.conn, self.book)
        self.assertEqual(r['machines_new'], 4)
        self.assertEqual(r['machines_updated'], 1)
        m = self.one("SELECT * FROM machines WHERE garage_no='5/104'")
        self.assertEqual((m['model'], m['plate'], m['reg_reading']), ('SHACMAN X3000', 'C000TO00', 47415))
        self.assertEqual(m['note'], 'моя заметка')                # в таблице пусто — в базе осталось
        self.assertEqual((m['reg_date'], m['norm_override']), ('2026-01-01', 90))
        self.assertEqual(self.one("SELECT COUNT(*) FROM machines WHERE garage_no LIKE '5/104%'")[0], 1)

    def test_rollback_on_error(self):
        with mock.patch.object(xlsx_io._Import, 'apply_units', side_effect=RuntimeError('сбой')):
            with self.assertRaises(RuntimeError):
                xlsx_io.import_v8(self.conn, self.book)
        self.assertFalse(self.conn.in_transaction)
        self.assertEqual(self.one('SELECT COUNT(*) FROM machines')[0], 0)
        self.assertEqual(self.one('SELECT COUNT(*) FROM events')[0], 0)
        self.assertEqual(db.get_settings(self.conn)['default_interval_mh'], 250)

    def test_bad_files(self):
        for data in (b'', b'not an xlsx at all', None):
            with self.assertRaises(ValueError) as cm:
                xlsx_io.import_v8(self.conn, data)
            self.assertIn('не книга Excel', str(cm.exception))

        wb = openpyxl.Workbook()
        wb.active.title = 'Машины'                      # листа «Журнал» нет
        buf = io.BytesIO()
        wb.save(buf)
        with self.assertRaises(ValueError) as cm:
            xlsx_io.import_v8(self.conn, buf.getvalue())
        self.assertIn('Это не таблица учёта ТО версии 8', str(cm.exception))
        self.assertFalse(self.conn.in_transaction)

    def test_only_required_sheets(self):
        """Достаточно листов «Машины» и «Журнал»; настройки и регламент при этом не меняются."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Машины'
        put(ws, 5, {'B': 'КАМАЗ', 'C': '77'})
        ws = wb.create_sheet('Журнал')
        put(ws, 5, {'A': datetime(2026, 7, 1), 'C': 77, 'D': 'Заметка'})        # гаражный № числом
        buf = io.BytesIO()
        wb.save(buf)
        r = xlsx_io.import_v8(self.conn, buf.getvalue())
        self.assertEqual((r['machines_new'], r['events_new'], r['settings_updated']), (1, 1, 0))
        self.assertEqual(r['warnings'], [])
        self.assertEqual(len(db.get_regulation(self.conn)), 3)

    def test_header_check(self):
        """Таблица другой версии (столбцы сдвинуты) — предупреждение о шапке; правильная шапка — без него."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Машины'
        put(ws, 4, {'A': '№', 'B': 'Модель', 'C': 'Гаражный №', 'D': 'Госномер', 'E': 'VIN',
                    'F': 'Гарантийный срок', 'K': 'Статус машины'})
        put(ws, 5, {'B': 'КАМАЗ', 'C': '77', 'F': 'до 12.2027'})
        ws = wb.create_sheet('Журнал')
        put(ws, 4, {'A': 'ДАТА *', 'C': 'МАШИНА *', 'D': 'СОБЫТИЕ *', 'G': 'Вышла — ДАТА',
                    'J': 'Что делалось / что сломалось'})
        buf = io.BytesIO()
        wb.save(buf)
        r = xlsx_io.import_v8(self.conn, buf.getvalue(), dry_run=True)
        self.assertWarning(r, 'Лист «Машины»', 'шапка', 'столбец F', 'Гарантийный срок')
        self.assertFalse(any('Лист «Журнал»' in w for w in r['warnings']))
        self.assertEqual(r['machines_new'], 1)                      # шапка — не машина

        r = xlsx_io.import_v8(self.conn, self.book, dry_run=True)
        self.assertFalse(any('шапка' in w for w in r['warnings']))

    def test_warnings_limit(self):
        """Предупреждений в отчёте не больше MAX_WARNINGS + строка-итог."""
        wb = openpyxl.Workbook()
        wb.active.title = 'Машины'
        ws = wb.create_sheet('Журнал')
        for r in range(5, 5 + xlsx_io.MAX_WARNINGS + 50):
            put(ws, r, {'A': datetime(2026, 7, 1), 'D': 'Заметка'})             # нет машины
        buf = io.BytesIO()
        wb.save(buf)
        r = xlsx_io.import_v8(self.conn, buf.getvalue())
        self.assertEqual(r['warnings_total'], xlsx_io.MAX_WARNINGS + 50)
        self.assertEqual(len(r['warnings']), xlsx_io.MAX_WARNINGS + 1)
        self.assertIn('ещё 50', r['warnings'][-1])


def make_book(machines=(), journal=(), tyres=(), units=()):
    """Небольшая книга v8: списки строк {буква: значение} для листов (данные с 5-й строки) -> bytes."""
    wb = openpyxl.Workbook()
    wb.active.title = 'Машины'
    for sheet, rows in (('Машины', machines), ('Журнал', journal), ('Шины', tyres), ('Узлы', units)):
        ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet)
        for i, row in enumerate(rows):
            put(ws, 5 + i, row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ImportCustomerDataTest(unittest.TestCase):
    """Случаи из реальной таблицы заказчика (v10): марка вместо № шины, дозаполнение журнала, латиница в номере."""

    def setUp(self):
        self.conn = db.open_db(':memory:')
        self.addCleanup(self.conn.close)

    def count(self, sql, *args):
        return self.conn.execute(sql, args).fetchone()[0]

    def test_tyres_brand_instead_of_number(self):
        """В «№ шины» — марка («СИТ»), разные шины отличаются серийным № в примечании: все импортируются."""
        machines = [{'B': 'SITRAK', 'C': '12А/410'}, {'B': 'SITRAK', 'C': '8А/225'}]
        tyres = [
            {'A': 'СИТ', 'B': '12А/410', 'C': '12.00R20', 'D': 'Перед. левая', 'K': 'с/н 0000000101'},
            {'A': 'СИТ', 'B': '12А/410', 'C': '12.00R20', 'D': 'Перед. левая', 'K': 'с/н 0000000102'},
            {'A': 'ШАХ', 'B': '8А/225', 'C': '20.5-25', 'D': 'Перед. правая'},
            {'A': 'ШАХ ', 'B': '8А/225', 'C': '23.5-25', 'D': 'Перед. правая'},
        ]
        units = [{'A': '8А/225', 'B': 'Гидроцилиндр', 'C': 3000, 'I': 'левый'},
                 {'A': '8А/225', 'B': 'Гидроцилиндр', 'C': 3000, 'I': 'правый'}]
        book = make_book(machines, tyres=tyres, units=units)
        r = xlsx_io.import_v8(self.conn, book)
        self.assertEqual((r['tyres_new'], r['tyres_skipped'], r['units_new']), (4, 0, 2))
        self.assertEqual(self.count('SELECT COUNT(*) FROM tyres'), 4)
        self.assertEqual(self.count('SELECT COUNT(*) FROM units'), 2)
        # повторный импорт того же файла — ничего не задваивается
        r = xlsx_io.import_v8(self.conn, book)
        self.assertEqual((r['tyres_new'], r['tyres_skipped'], r['units_new'], r['units_skipped']), (0, 4, 0, 2))
        self.assertEqual(self.count('SELECT COUNT(*) FROM tyres'), 4)

    def test_repeat_import_fills_journal(self):
        """Повторный импорт: дописанные в таблице «вышла»/замечания попадают в базу; противоречие — предупреждение."""
        machines = [{'B': 'SHACMAN', 'C': '14А/508'}]
        repair = {'A': datetime(2026, 8, 16), 'C': '14А/508', 'D': 'Ремонт', 'J': 'замена кабины'}
        note = {'A': datetime(2026, 8, 20), 'C': '14А/508', 'D': 'Заметка', 'M': 'шум в мосту'}
        xlsx_io.import_v8(self.conn, make_book(machines, [repair, note]))

        closed = dict(repair, G=datetime(2026, 9, 5), H=time(18, 0), K='кабина')
        note2 = dict(note, M='течь ГУР')                  # другая заметка за тот же день — это новая запись
        r = xlsx_io.import_v8(self.conn, make_book(machines, [closed, note, note2]))
        self.assertEqual((r['events_new'], r['events_updated'], r['events_skipped']), (1, 1, 1))
        ev = self.conn.execute("SELECT * FROM events WHERE kind='Ремонт'").fetchone()
        self.assertEqual((ev['date_out'], ev['time_out'], ev['parts_used']), ('2026-09-05', '18:00', 'кабина'))
        self.assertEqual(self.count("SELECT COUNT(*) FROM events WHERE kind='Заметка'"), 2)

        # в таблице другая дата выхода — база не меняется, предупреждение с номером строки
        other = dict(closed, G=datetime(2026, 9, 6))
        r = xlsx_io.import_v8(self.conn, make_book(machines, [other]))
        self.assertEqual((r['events_new'], r['events_updated'], r['events_skipped']), (0, 0, 1))
        self.assertTrue(any('Журнал, строка 5' in w and 'дата выхода' in w for w in r['warnings']), r['warnings'])
        ev = self.conn.execute("SELECT * FROM events WHERE kind='Ремонт'").fetchone()
        self.assertEqual(ev['date_out'], '2026-09-05')

    def test_latin_garage_no_fixed_in_program(self):
        """«2A/317» латиницей исправлен в программе на кириллицу — повторный импорт не создаёт машину-дубль."""
        machines = [{'B': 'SITRAK', 'C': '2A/317'}]
        journal = [{'A': datetime(2026, 8, 18), 'C': '2A/317', 'D': 'Ремонт', 'J': 'стартер'}]
        book = make_book(machines, journal)
        r = xlsx_io.import_v8(self.conn, book)
        self.assertTrue(any('2A/317' in w and 'латинские буквы' in w for w in r['warnings']), r['warnings'])
        self.conn.execute("UPDATE machines SET garage_no='2А/317'")
        self.conn.commit()
        r = xlsx_io.import_v8(self.conn, book)
        self.assertEqual((r['machines_new'], r['events_new'], r['events_skipped']), (0, 0, 1))
        self.assertTrue(any('2A/317' in w and '2А/317' in w for w in r['warnings']), r['warnings'])
        self.assertEqual(self.count('SELECT COUNT(*) FROM machines'), 1)
        self.assertEqual(self.count('SELECT COUNT(*) FROM events'), 1)


class ExportTest(unittest.TestCase):
    """Экспорт и обратный импорт."""

    def setUp(self):
        self.conn = db.open_db(':memory:')
        self.addCleanup(self.conn.close)
        c = self.conn
        now = db.now_iso()
        c.execute("INSERT INTO machines(id, pos, garage_no, model, mtype, plate, vin, site, warranty, note, status, "
                  "reg_reading, reg_date, created_at, updated_at) VALUES(1, 1, '5/104', 'SHACMAN X3000', 'Самосвал', "
                  "'C000TO00', 'TESTVIN0000000001', 'Северный', 'до 2027', 'основная', 'В работе', 47415, NULL, ?, ?)",
                  (now, now))
        c.execute("INSERT INTO machines(id, pos, garage_no, model, mtype, status, created_at, updated_at) "
                  "VALUES(2, 2, '4Э', 'Экскаватор SANY 330', 'Экскаватор', 'Списана', ?, ?)", (now, now))
        events = [
            # machine, date, shift, kind, reading, estimated, delta, source, time_in, date_out, time_out,
            # work, parts_used, parts_order, order_done, remarks, who
            (1, '2026-09-01', 'День', 'ТО-1', 150000, 0, None, 'manual', None, None, None,
             'Замена масла', 'масло 10 л', '', 0, '', 'Петров'),
            (1, '2026-09-10', '', KIND_MARK, 153000, 0, None, 'mark', None, None, None, '', '', '', 0, '', ''),
            (1, '2026-09-11', '', KIND_MARK, None, 0, 100, 'auto', None, None, None, '', '', '', 0, '', 'авто'),
            (1, '2026-09-12', '', KIND_MARK, None, 0, 100, 'auto', None, None, None, '', '', '', 0, '', 'авто'),
            (1, '2026-09-13', 'Ночь', 'Ремонт', 153200, 0, None, 'manual', '09:30', '2026-09-14', '13:50',
             'Замена стартера', 'стартер', 'стартер SHACMAN', 1, 'стружка в масле', 'Иванов'),
            (1, '2026-09-20', 'День', 'Ремонт', None, 0, None, 'manual', '08:00', None, None,
             '=не формула', '', 'прокладка ГБЦ', 0, '', ''),
            (2, '2026-09-05', '', KIND_MARK, 4372.5, 0, None, 'mark', None, None, None, '', '', '', 0, '', ''),
            (2, '2026-09-06', '', KIND_MARK, None, 0, 11.5, 'auto', None, None, None, '', '', '', 0, '', 'авто'),
        ]
        c.executemany(
            'INSERT INTO events(machine_id, date, shift, kind, reading, reading_estimated, delta, source, time_in, '
            'date_out, time_out, work, parts_used, parts_order, order_done, remarks, who, created_at, updated_at) '
            f"VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, '{now}', '{now}')", events)
        c.execute("INSERT INTO tyres(machine_id, tyre_no, ttype, position, install_date, resource, run, note) "
                  "VALUES(1, '22', '12.00R20', 'Перед. правая', '2026-06-30', 80000, 76000, 'пора менять')")
        c.execute("INSERT INTO units(machine_id, unit, resource, install_reading, note) "
                  "VALUES(1, 'КПП', 200000, 100000, 'масло 80W-90')")
        c.commit()
        self.n_events = len(events)

    def export(self):
        data = xlsx_io.export_xlsx(self.conn)
        self.assertIsInstance(data, bytes)
        return data

    def test_export_opens(self):
        wb = openpyxl.load_workbook(io.BytesIO(self.export()))
        for name in ('Машины', 'Журнал', 'Шины', 'Узлы'):
            self.assertIn(name, wb.sheetnames)
        expected_headers = {
            'Машины': ['№', 'Модель / тип техники', 'Гаражный №', 'Госномер', 'VIN'],
            'Журнал': ['ДАТА *', 'Смена', 'МАШИНА *', 'СОБЫТИЕ *', 'МОТОЧАСЫ / КМ *', 'Заехала — ВРЕМЯ',
                       'Вышла — ДАТА', 'Вышла — ВРЕМЯ', 'В ремонте ВСЕГО (само)'],
            'Шины': ['№ шины', 'Гаражный №', 'Тип / размер шины', 'Позиция', 'Дата установки'],
            'Узлы': ['Машина (гар. №)', 'Узел'],
        }
        for name, headers in expected_headers.items():
            ws = wb[name]
            self.assertTrue(ws['A1'].value, name)                     # заголовок листа в 1-й строке
            self.assertTrue(ws['A1'].font.bold, name)
            self.assertIsNone(ws['A3'].value, name)
            row4 = [cell.value for cell in ws[4]]
            self.assertEqual(row4[:len(headers)], headers, name)      # шапка в 4-й строке
            self.assertTrue(all(cell.font.bold for cell in ws[4]), name)
            self.assertEqual(ws.freeze_panes, 'A5', name)             # шапка закреплена
            self.assertTrue(ws.auto_filter.ref.startswith('A4:'), name)
            # значения без формул
            formulas = [cell.coordinate for row in ws.iter_rows() for cell in row if cell.data_type == 'f']
            self.assertEqual(formulas, [], name)

        ws = wb['Журнал']
        self.assertEqual(ws.max_row, 4 + self.n_events)
        # A–N как в таблице v8, O–P — служебные столбцы выгрузки («Заказано», «Служебное»)
        self.assertEqual(ws.auto_filter.ref, f'A4:P{4 + self.n_events}')
        rows = {(row[0].date().isoformat(), row[2]): row for row in ws.iter_rows(min_row=5, values_only=True)}
        ev = rows[('2026-09-13', '5/104')]
        self.assertEqual(ev[3:9], ('Ремонт', 153200, time(9, 30), datetime(2026, 9, 14), time(13, 50), '1 д 4 ч 20 мин'))
        self.assertEqual(ev[9:14], ('Замена стартера', 'стартер', 'стартер SHACMAN', 'стружка в масле', 'Иванов'))
        auto = rows[('2026-09-11', '5/104')]
        self.assertEqual((auto[3], auto[4], auto[13]), (KIND_MARK, None, 'авто +100'))
        self.assertEqual(rows[('2026-09-06', '4Э')][13], 'авто +11.5')
        self.assertEqual(rows[('2026-09-20', '5/104')][9], ' =не формула')     # текст (с пробелом), а не формула
        self.assertIsNone(rows[('2026-09-20', '5/104')][8])                    # открытый ремонт — итога нет

        ws = wb['Машины']
        self.assertEqual([c.value for c in ws[5]][:12], [1, 'SHACMAN X3000', '5/104', 'C000TO00',
                                                         'TESTVIN0000000001', 47415, 'Северный', 'до 2027', None,
                                                         'основная', 'В работе', 'Самосвал'])
        self.assertEqual(ws['K6'].value, 'Списана')                           # списанные тоже выгружаются

        ws = wb['Шины']
        self.assertEqual([c.value for c in ws[5]][:7] + [ws['J5'].value, ws['K5'].value],
                         ['22', '5/104', '12.00R20', 'Перед. правая', datetime(2026, 6, 30), 80000, 76000,
                          '🔴 Менять', 'пора менять'])
        self.assertEqual((ws['H5'].value, ws['I5'].value), (4000, 0.95))
        ws = wb['Узлы']
        self.assertEqual([c.value for c in ws[5]][:4] + [ws['I5'].value],
                         ['5/104', 'КПП', 200000, 100000, 'масло 80W-90'])

    def test_export_empty_db(self):
        conn = db.open_db(':memory:')
        self.addCleanup(conn.close)
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_io.export_xlsx(conn)))
        self.assertEqual(wb['Журнал'].max_row, 4)
        self.assertEqual(wb['Журнал']['A4'].value, 'ДАТА *')
        # пустая выгрузка тоже импортируется
        r = xlsx_io.import_v8(conn, xlsx_io.export_xlsx(conn))
        self.assertEqual((r['machines_new'], r['events_new'], r['warnings']), (0, 0, []))

    def test_export_then_import_into_new_db(self):
        data = self.export()
        new = db.open_db(':memory:')
        self.addCleanup(new.close)
        r = xlsx_io.import_v8(new, data)
        self.assertEqual(r['machines_new'], 2)
        self.assertEqual(r['machines_updated'], 0)
        self.assertEqual(r['events_new'], self.n_events)
        self.assertEqual(r['events_skipped'], 0)
        self.assertEqual((r['tyres_new'], r['units_new']), (1, 1))
        self.assertEqual(r['orders_open'], 1)
        self.assertEqual(r['settings_updated'], 0)
        self.assertEqual(r['warnings'], [])

        def machines(conn):
            return [tuple(row) for row in conn.execute(
                'SELECT garage_no, model, mtype, plate, vin, site, warranty, note, status, reg_reading '
                'FROM machines ORDER BY pos, id')]

        def events(conn):
            return [tuple(row) for row in conn.execute(
                'SELECT m.garage_no, e.date, e.shift, e.kind, e.reading, e.delta, e.time_in, e.date_out, e.time_out, '
                'e.work, e.parts_used, e.parts_order, e.order_done, e.remarks, e.who '
                'FROM events e JOIN machines m ON m.id=e.machine_id ORDER BY e.date, e.id')]

        self.assertEqual(machines(new), machines(self.conn))
        self.assertEqual(events(new), events(self.conn))
        # текст, начинающийся с «=», не потерялся и не стал формулой
        self.assertEqual(new.execute("SELECT work FROM events WHERE date='2026-09-20'").fetchone()[0], '=не формула')

        # накинутое восстановилось как auto
        autos = [tuple(row) for row in new.execute(
            "SELECT date, reading, delta, who, kind FROM events WHERE source='auto' ORDER BY date")]
        self.assertEqual(autos, [('2026-09-06', None, 11.5, 'авто', KIND_MARK),
                                 ('2026-09-11', None, 100, 'авто', KIND_MARK),
                                 ('2026-09-12', None, 100, 'авто', KIND_MARK)])
        # отметки с экрана «Показания» остаются отметками (служебный столбец P), остальное — 'import'
        self.assertEqual([tuple(r) for r in new.execute(
            "SELECT date, source FROM events WHERE source NOT IN ('auto', 'import') ORDER BY date")],
            [('2026-09-05', 'mark'), ('2026-09-10', 'mark')])

        tyre = new.execute('SELECT tyre_no, ttype, position, install_date, resource, run, note FROM tyres').fetchone()
        self.assertEqual(tuple(tyre), ('22', '12.00R20', 'Перед. правая', '2026-06-30', 80000, 76000, 'пора менять'))
        unit = new.execute('SELECT unit, resource, install_reading, note FROM units').fetchone()
        self.assertEqual(tuple(unit), ('КПП', 200000, 100000, 'масло 80W-90'))

        # повторный импорт той же выгрузки — без дублей
        r = xlsx_io.import_v8(new, data)
        self.assertEqual((r['machines_new'], r['machines_updated'], r['events_new'], r['tyres_new'], r['units_new']),
                         (0, 0, 0, 0, 0))
        self.assertEqual(r['events_skipped'], self.n_events)

    def test_export_then_import_into_same_db(self):
        """Выгрузку можно загрузить обратно в ту же базу: ничего не дублируется и не меняется."""
        before = [tuple(row) for row in self.conn.execute('SELECT * FROM events ORDER BY id')]
        r = xlsx_io.import_v8(self.conn, self.export())
        self.assertEqual((r['machines_new'], r['machines_updated'], r['events_new'], r['tyres_new'], r['units_new']),
                         (0, 0, 0, 0, 0))
        self.assertEqual(r['events_skipped'], self.n_events)
        self.assertEqual(before, [tuple(row) for row in self.conn.execute('SELECT * FROM events ORDER BY id')])

    def test_roundtrip_keeps_service_fields(self):
        """Интеграция: выгрузка -> чистая база без потерь того, чего нет в столбцах v8:
        baseline-ТО, «≈» у ТО, открытая заявка у закрытого ремонта, свои параметры ТО и reg_date машины,
        показание при установке шины."""
        c = self.conn
        now = db.now_iso()
        c.execute("UPDATE machines SET unit_override='км', interval_override=12000, warn_override=800, "
                  "norm_override=150, reg_date='2026-08-01' WHERE id=1")
        c.execute("UPDATE events SET order_done=0 WHERE date='2026-09-13'")      # закрытый ремонт, не заказано
        c.execute("INSERT INTO events(machine_id, date, kind, reading, source, created_at, updated_at) "
                  "VALUES(1, '2026-09-15', 'ТО-2', 153000, 'baseline', ?, ?)", (now, now))
        c.execute("INSERT INTO events(machine_id, date, kind, reading, reading_estimated, source, created_at, "
                  "updated_at) VALUES(1, '2026-09-16', 'ТО-3', 153400, 1, 'manual', ?, ?)", (now, now))
        c.execute("UPDATE tyres SET install_reading=100000")
        c.commit()

        new = db.open_db(':memory:')
        self.addCleanup(new.close)
        r = xlsx_io.import_v8(new, self.export())
        self.assertEqual(r['warnings'], [])
        self.assertEqual(r['orders_open'], 2)                    # открытый ремонт + закрытый с незаказанной заявкой

        cols = 'unit_override, interval_override, warn_override, norm_override, reg_date'
        self.assertEqual([tuple(x) for x in new.execute(f'SELECT {cols} FROM machines ORDER BY pos')],
                         [tuple(x) for x in c.execute(f'SELECT {cols} FROM machines ORDER BY pos')])
        q = ('SELECT e.date, e.kind, e.reading, e.reading_estimated, e.order_done, e.source FROM events e '
             'ORDER BY e.date, e.id')
        norm = {'manual': 'import'}                              # ручные записи после импорта — 'import'
        self.assertEqual([tuple(x) for x in new.execute(q)],
                         [tuple(x)[:5] + (norm.get(x['source'], x['source']),) for x in c.execute(q)])
        self.assertEqual(new.execute('SELECT install_reading FROM tyres').fetchone()[0], 100000)

        # повторный импорт той же выгрузки в эту же базу — без дублей и без ошибки уникальности отметок
        r = xlsx_io.import_v8(new, self.export())
        self.assertEqual(r['events_new'], 0)

    def test_roundtrip_keeps_date_unknown(self):
        """«Последнее ТО» без даты (date_unknown=1): в служебном столбце P — «baseline;без даты»,
        на листе «Парк» вместо даты ТО — «дата не указана»; обратный импорт восстанавливает признак."""
        c = self.conn
        now = db.now_iso()
        c.execute("INSERT INTO events(machine_id, date, kind, reading, source, date_unknown, created_at, updated_at) "
                  "VALUES(1, '2026-09-15', 'ТО-2', 153000, 'baseline', 1, ?, ?)", (now, now))
        c.commit()
        data = self.export()
        wb = openpyxl.load_workbook(io.BytesIO(data))
        rows = {(r[0].date().isoformat(), r[2], r[3]): r for r in wb['Журнал'].iter_rows(min_row=5, values_only=True)}
        self.assertEqual(rows[('2026-09-15', '5/104', 'ТО-2')][15], 'baseline;без даты')
        self.assertIsNone(rows[('2026-09-01', '5/104', 'ТО-1')][15])
        if 'Парк' in wb.sheetnames:
            self.assertEqual(wb['Парк']['L5'].value, 'дата не указана')

        new = db.open_db(':memory:')
        self.addCleanup(new.close)
        r = xlsx_io.import_v8(new, data)
        self.assertEqual(r['warnings'], [])
        q = "SELECT date, kind, source, date_unknown FROM events WHERE kind LIKE 'ТО-%' ORDER BY date"
        self.assertEqual([tuple(x) for x in new.execute(q)],
                         [('2026-09-01', 'ТО-1', 'import', 0), ('2026-09-15', 'ТО-2', 'baseline', 1)])

    def test_mark_on_same_date_imported_without_conflict(self):
        """Отметка из выгрузки на дату, где в базе уже есть другая отметка (source='mark') с другим показанием,
        импортируется как 'import' — без ошибки уникального индекса ux_ev_mark."""
        data = self.export()
        self.conn.execute("UPDATE events SET reading=153050 WHERE source='mark' AND date='2026-09-10'")
        self.conn.commit()
        r = xlsx_io.import_v8(self.conn, data)
        self.assertEqual(r['events_new'], 1)
        self.assertEqual([tuple(x) for x in self.conn.execute(
            "SELECT reading, source FROM events WHERE date='2026-09-10' ORDER BY id")],
            [(153050, 'mark'), (153000, 'import')])

    def test_v8_ordered_column(self):
        """Столбец O «Заказано (дата)» таблицы v8: заполнен — заявка заказана (даже у открытого ремонта)."""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet('Журнал')
        put(ws, 4, dict(zip('ABCDEFGHIJKLMNO', [
            'ДАТА *', 'Смена', 'МАШИНА *', 'СОБЫТИЕ *', 'МОТОЧАСЫ / КМ *', 'Заехала — ВРЕМЯ', 'Вышла — ДАТА',
            'Вышла — ВРЕМЯ', 'В ремонте ВСЕГО (само)', 'Что делалось / что сломалось',
            'Масло / запчасти (использовано)', 'Заказать запчасти (в закуп)', 'Замечания (стружка и т.п.)', 'Кто',
            'Заказано (дата)'])))
        put(ws, 5, {'A': datetime(2026, 7, 1), 'C': 'X1', 'D': 'Ремонт', 'L': 'прокладка', 'O': datetime(2026, 7, 2)})
        put(ws, 6, {'A': datetime(2026, 7, 1), 'C': 'X1', 'D': 'Ремонт', 'L': 'сальник'})
        put(ws, 7, {'A': datetime(2026, 7, 1), 'C': 'X1', 'D': 'Заметка', 'L': 'фильтр'})
        wb.create_sheet('Машины')['C5'] = 'X1'
        buf = io.BytesIO()
        wb.save(buf)
        conn = db.open_db(':memory:')
        self.addCleanup(conn.close)
        r = xlsx_io.import_v8(conn, buf.getvalue())
        self.assertEqual(r['orders_open'], 1)
        self.assertEqual([tuple(x) for x in conn.execute('SELECT parts_order, order_done FROM events ORDER BY id')],
                         [('прокладка', 1), ('сальник', 0), ('фильтр', 1)])

    @unittest.skipUnless(services_available(), 'app/services.py (build_views) ещё не готов — лист «Парк» не строится')
    def test_export_fleet_sheet(self):
        wb = openpyxl.load_workbook(io.BytesIO(self.export()))
        self.assertEqual(wb.sheetnames, ['Парк', 'Машины', 'Журнал', 'Шины', 'Узлы'])
        ws = wb['Парк']
        self.assertIn('Парк', ws['A1'].value)
        self.assertEqual([c.value for c in ws[4]][:4], ['№', 'Гаражный №', 'Модель', 'Участок'])
        self.assertEqual(ws.freeze_panes, 'A5')
        self.assertEqual(ws.max_row, 5)                                # списанная машина в «Парк» не попадает
        row = [c.value for c in ws[5]]
        self.assertEqual(row[:3], [1, '5/104', 'SHACMAN X3000'])
        self.assertEqual(row[6:9], [153200, None, 'км'])               # показание: ремонт 153200 выше отметки
        self.assertEqual(row[9:11], ['ТО-1', 150000])

    def test_export_without_services(self):
        """Если services недоступен, экспорт работает без листа «Парк»."""
        with mock.patch.object(xlsx_io, '_fleet_views', return_value=None):
            wb = openpyxl.load_workbook(io.BytesIO(self.export()))
        self.assertEqual(wb.sheetnames, ['Машины', 'Журнал', 'Шины', 'Узлы'])


if __name__ == '__main__':
    unittest.main()
